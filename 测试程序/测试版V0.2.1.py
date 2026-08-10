import time
print("=" * 50)
print("本程序由繁星攻略组制作")
print("QQ：1035835099")
print("WX：TangTangMei18")
print("=" * 50)
print(" ")

time.sleep(5)  # 延时5秒

import openpyxl
from openpyxl import load_workbook
import os
import sys
from decimal import Decimal, getcontext, ROUND_HALF_UP

class ArmorPenetrationCalculator:
    def __init__(self):
        # ---------- Debug 模式检测 ----------
        self.debug = False
        # 方式1: 代码中直接设为 True (永久开启)
        # self.debug = True
        
        # 方式2: 检测脚本文件名是否包含 ".debug" (例如: calculator.debug.py)
        if '.debug' in sys.argv[0]:
            self.debug = True
        
        if self.debug:
            print("[DEBUG] 调试模式已启用，将输出详细计算信息。")
        
        # 设置Decimal精度
        getcontext().prec = 10
        
        print("正在加载数据文件...")
        
        # 加载穿透数据
        self.penetration_data = {}
        self.load_penetration_data()
        
        # 加载武器数据
        self.weapons_by_category = {}
        self.weapons_data = {}
        self.load_weapon_data()
        
        # 加载弹药数据
        self.ammo_by_caliber = {}
        self.ammo_data = {}
        self.load_ammo_data()
        
        # 加载护甲数据
        self.armor_by_level = {}
        self.armor_data = {}
        self.load_armor_data()
        
        print("数据加载完成！\n")
    
    def standardize_caliber(self, caliber_str):
        """标准化口径字符串，使其能够正确匹配"""
        if not caliber_str:
            return ""
        caliber = str(caliber_str).strip()
        caliber = caliber.upper()
        caliber = caliber.replace('X', 'x')
        caliber = caliber.replace(' ', '')
        caliber = caliber.replace('MM', 'mm')
        
        # 处理特定格式
        if caliber == '5.56X45MM':
            caliber = '5.56x45mm'
        elif caliber == '5.45X39MM':
            caliber = '5.45x39mm'
        elif caliber == '7.62X39MM':
            caliber = '7.62x39mm'
        elif caliber == '7.62X51MM':
            caliber = '7.62x51mm'
        elif caliber == '7.62X54MM':
            caliber = '7.62x54mm'
        elif caliber == '5.8X42MM':
            caliber = '5.8x42mm'
        elif caliber == '9X19MM':
            caliber = '9x19mm'
        elif caliber == '9X39MM':
            caliber = '9x39mm'
        elif caliber == '12 GAUGE':
            caliber = '12 Gauge'
        elif caliber == '.45 ACP':
            caliber = '.45 ACP'
        elif caliber == '.44 MAG':
            caliber = '.44 Mag'
        elif caliber == '.338 LAPUA MAG':
            caliber = '.338 Lapua Mag'
        return caliber
    
    def load_penetration_data(self):
        """加载穿透百分比数据"""
        try:
            wb = load_workbook('穿透数据整理结果.xlsx')
            ws = wb.active
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] is not None and row[3] is not None:
                    k_value = int(row[0])
                    self.penetration_data[k_value] = Decimal(str(row[3]))
            print(f"已加载 {len(self.penetration_data)} 条穿透数据")
            if self.debug:
                print("\n[DEBUG] 穿透数据映射表 (K值 → 穿透百分比):")
                for k, v in sorted(self.penetration_data.items()):
                    print(f"  K={k:3d}  →  {v:.4f}")
                print()
        except Exception as e:
            print(f"加载穿透数据失败: {e}")
            self.penetration_data = {
                -10: Decimal('0.6'), -9: Decimal('0.6053'), -8: Decimal('0.611'),
                -7: Decimal('0.615'), -6: Decimal('0.62'), -5: Decimal('0.626'),
                -4: Decimal('0.628'), -3: Decimal('0.639'), -1: Decimal('0.645'),
                0: Decimal('0.65'), 1: Decimal('0.6556'), 2: Decimal('0.66'),
                3: Decimal('0.666'), 4: Decimal('0.67'), 5: Decimal('0.675'),
                6: Decimal('0.681'), 7: Decimal('0.685'), 8: Decimal('0.688'),
                9: Decimal('0.696'), 10: Decimal('0.9'), 11: Decimal('0.91'),
                12: Decimal('0.92'), 13: Decimal('0.93'), 14: Decimal('0.94'),
                15: Decimal('0.95'), 17: Decimal('0.97'), 18: Decimal('0.98')
            }
            if self.debug:
                print("[DEBUG] 使用内置备选穿透数据。")
    
    def load_weapon_data(self):
        """加载武器数据"""
        try:
            wb = load_workbook('武器数据.xlsx')
            ws = wb.active
            current_category = None
            weapon_id = 0
            for row in ws.iter_rows(min_row=3, values_only=True):
                if row[0] and not row[1]:
                    current_category = row[0]
                    self.weapons_by_category[current_category] = []
                elif row[0] and row[1]:
                    weapon_id += 1
                    caliber = self.standardize_caliber(row[1])
                    weapon = {
                        'id': weapon_id,
                        'name': row[0],
                        'category': current_category,
                        'caliber': caliber,
                        'original_caliber': row[1],
                        'barrels': []
                    }
                    barrel_name = row[4] if row[4] else "/"
                    barrel = {
                        'name': barrel_name,
                        'firing_power': Decimal(str(row[5])) if row[5] is not None else Decimal('0')
                    }
                    weapon['barrels'].append(barrel)
                    self.weapons_by_category[current_category].append(weapon)
                    self.weapons_data[weapon_id] = weapon
                elif not row[0] and row[4]:
                    weapon = self.weapons_by_category[current_category][-1]
                    barrel = {
                        'name': row[4],
                        'firing_power': Decimal(str(row[5])) if row[5] is not None else Decimal('0')
                    }
                    weapon['barrels'].append(barrel)
            print(f"已加载 {len(self.weapons_data)} 种武器，分为 {len(self.weapons_by_category)} 个类别")
            if self.debug:
                print("[DEBUG] 武器加载详情（仅显示前3个）:")
                for cat, weapons in list(self.weapons_by_category.items())[:3]:
                    print(f"  类别: {cat}")
                    for w in weapons[:3]:
                        print(f"    {w['name']} (口径: {w['original_caliber']}) 枪管数: {len(w['barrels'])}")
                print()
        except Exception as e:
            print(f"加载武器数据失败: {e}")
    
    def load_ammo_data(self):
        """加载弹药数据 - 修正列索引"""
        try:
            wb = load_workbook('子弹数据.xlsx')
            ws = wb.active
            ammo_id = 0
            current_caliber = None
            current_original_caliber = None
            
            for row in ws.iter_rows(min_row=3, values_only=True):
                if row[0]:  # A列：口径
                    current_original_caliber = row[0]
                    current_caliber = self.standardize_caliber(row[0])
                    if current_caliber not in self.ammo_by_caliber:
                        self.ammo_by_caliber[current_caliber] = []
                
                if row[1]:  # B列：弹药型号
                    ammo_id += 1
                    
                    # 穿透等级：F列（索引5）
                    penetration_level = 0
                    if row[5] is not None:
                        try:
                            penetration_level = int(float(row[5]))
                        except:
                            pass
                    
                    # 子弹穿透值：G列（索引6）
                    bullet_penetration = Decimal('0')
                    if row[6] is not None:
                        try:
                            bullet_penetration = Decimal(str(row[6]))
                        except:
                            bullet_penetration = Decimal('0')
                    
                    ammo = {
                        'id': ammo_id,
                        'caliber': current_caliber,
                        'original_caliber': current_original_caliber if current_original_caliber else current_caliber,
                        'name': row[1],
                        'damage': self.parse_damage_value(row[2]),       # C列：子弹伤害
                        'armor_damage': Decimal(str(row[3])) if row[3] else Decimal('0'),  # D列：护甲伤害
                        'penetration_level': penetration_level,
                        'bullet_penetration': bullet_penetration
                    }
                    
                    # 如果穿透等级为0，根据子弹穿透值计算
                    if ammo['penetration_level'] == 0 and ammo['bullet_penetration'] > 0:
                        ammo['penetration_level'] = int(ammo['bullet_penetration'] // 10)
                    
                    self.ammo_by_caliber[current_caliber].append(ammo)
                    self.ammo_data[ammo_id] = ammo
            
            print(f"已加载 {len(self.ammo_data)} 种子弹，涵盖 {len(self.ammo_by_caliber)} 种口径")
            print("可用口径:", ", ".join(sorted(self.ammo_by_caliber.keys())))
            if self.debug:
                print("[DEBUG] 弹药加载详情（每个口径前2种子弹）:")
                for cal, ammos in list(self.ammo_by_caliber.items())[:5]:
                    print(f"  口径: {cal}")
                    for a in ammos[:2]:
                        print(f"    {a['name']} 伤害:{a['damage']} 穿甲等级:{a['penetration_level']} 子弹穿透:{a['bullet_penetration']}")
                print()
        except Exception as e:
            print(f"加载弹药数据失败: {e}")
    
    def parse_damage_value(self, damage_str):
        """解析伤害值（处理像'86*2'这样的格式）"""
        if not damage_str:
            return Decimal('0')
        damage_str = str(damage_str)
        if '*' in damage_str:
            parts = damage_str.split('*')
            try:
                return Decimal(str(parts[0])) * Decimal(str(parts[1]))
            except:
                return Decimal('0')
        else:
            try:
                return Decimal(str(damage_str))
            except:
                return Decimal('0')
    
    def load_armor_data(self):
        """加载护甲数据"""
        try:
            wb = load_workbook('护甲信息.xlsx')
            ws = wb.active
            armor_id = 0
            for row in ws.iter_rows(min_row=3, values_only=True):
                if row[0] and row[1]:
                    armor_id += 1
                    material = row[2] if row[2] else "未知材质"
                    durability = Decimal('0')
                    if row[3] is not None:
                        try:
                            durability = Decimal(str(row[3]))
                        except:
                            pass
                    destruction_coef = Decimal('0')
                    if row[4] is not None:
                        try:
                            destruction_coef = Decimal(str(row[4]))
                        except:
                            pass
                    armor = {
                        'id': armor_id,
                        'name': row[0],
                        'level': int(row[1]),
                        'material': material,
                        'durability': durability,
                        'destruction_coef': destruction_coef,
                        'current_durability': durability
                    }
                    level = armor['level']
                    if level not in self.armor_by_level:
                        self.armor_by_level[level] = []
                    self.armor_by_level[level].append(armor)
                    self.armor_data[armor_id] = armor
            print(f"已加载 {len(self.armor_data)} 件护甲")
            if self.debug:
                print("[DEBUG] 护甲加载详情（每个等级前2件）:")
                for lvl, armors in list(self.armor_by_level.items())[:5]:
                    print(f"  等级 {lvl}:")
                    for a in armors[:2]:
                        print(f"    {a['name']} 耐久:{a['durability']} 损毁系数:{a['destruction_coef']}")
                print()
        except Exception as e:
            print(f"加载护甲数据失败: {e}")
    
    def check_penetration_status(self, ammo_penetration_level, armor_level, armor_current_durability, armor_max_durability):
        if ammo_penetration_level >= armor_level:
            return True, "必定击穿护甲"
        level_difference = armor_level - ammo_penetration_level
        if armor_max_durability <= Decimal('0'):
            return True, "必定击穿护甲"
        durability_ratio = armor_current_durability / armor_max_durability
        import math
        if durability_ratio <= Decimal('0'):
            return True, "必定击穿护甲"
        ratio_float = float(durability_ratio)
        if ratio_float > 0:
            half_lives_completed = math.floor(-math.log2(ratio_float))
        else:
            half_lives_completed = float('inf')
        if half_lives_completed >= level_difference:
            return True, "必定击穿护甲"
        elif half_lives_completed == level_difference - 1:
            return True, "概率击穿护甲"
        else:
            return False, "无法击穿护甲"
    
    def calculate_penetration_percentage(self, k_value):
        k_value_int = int(k_value)
        if self.debug:
            print(f"[DEBUG] calculate_penetration_percentage: k_value={k_value}, k_int={k_value_int}")
        if k_value_int < -10:
            result = Decimal('0.6')
            if self.debug:
                print("[DEBUG] K值 < -10，固定返回 0.6")
        elif k_value_int > 20:
            result = Decimal('1.0')
            if self.debug:
                print("[DEBUG] K值 > 20，固定返回 1.0")
        else:
            if k_value_int in self.penetration_data:
                result = self.penetration_data[k_value_int]
                if self.debug:
                    print(f"[DEBUG] 精确匹配 K={k_value_int}")
            else:
                closest_k = min(self.penetration_data.keys(), key=lambda x: abs(x - k_value_int))
                result = self.penetration_data[closest_k]
                if self.debug:
                    print(f"[DEBUG] 未精确匹配，使用最近 K={closest_k}")
        if self.debug:
            print(f"[DEBUG] 返回穿透百分比: {result:.4f}")
        return result
    
    def calculate_damage(self, weapon, barrel, ammo, armor, current_durability, current_max_durability=None):
        if current_max_durability is None:
            current_max_durability = armor['durability']
        if self.debug:
            print(f"[DEBUG] calculate_damage 输入: 弹药穿透等级={ammo['penetration_level']}, 护甲等级={armor['level']}, 当前耐久={current_durability}, 上限={current_max_durability}")
        can_penetrate, penetration_type = self.check_penetration_status(
            ammo['penetration_level'],
            armor['level'],
            current_durability,
            current_max_durability
        )
        armor_damage = (ammo['armor_damage'] * armor['destruction_coef']).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        new_durability = (current_durability - armor_damage).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        if new_durability < Decimal('0'):
            new_durability = Decimal('0')
        if self.debug:
            print(f"[DEBUG] 击穿状态: {penetration_type}, can_penetrate={can_penetrate}")
        if not can_penetrate:
            return {
                'can_penetrate': False,
                'penetration_type': penetration_type,
                'damage': Decimal('0'),
                'armor_damage_dealt': armor_damage,
                'remaining_durability': new_durability,
                'is_destroyed': new_durability <= Decimal('0'),
                'current_max_durability': current_max_durability
            }
        k_value = ammo['bullet_penetration'] - Decimal(str(armor['level'])) * Decimal('10')
        if self.debug:
            print(f"[DEBUG] K值计算: {ammo['bullet_penetration']} - {armor['level']}*10 = {k_value}")
        penetration_percentage = self.calculate_penetration_percentage(k_value)
        base_damage = ammo['damage'] + barrel['firing_power']
        damage = (base_damage * penetration_percentage).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
        if self.debug:
            print(f"[DEBUG] 基础伤害={base_damage}, 穿透百分比={penetration_percentage:.4f}, 最终伤害={damage}")
        return {
            'can_penetrate': True,
            'penetration_type': penetration_type,
            'k_value': k_value,
            'penetration_percentage': penetration_percentage,
            'damage': damage,
            'armor_damage_dealt': armor_damage,
            'remaining_durability': new_durability,
            'is_destroyed': new_durability <= Decimal('0'),
            'current_max_durability': current_max_durability
        }
    
    def select_weapon(self):
        print("=" * 50)
        print("请选择武器类别:")
        categories = list(self.weapons_by_category.keys())
        for i, category in enumerate(categories, 1):
            weapon_count = len(self.weapons_by_category[category])
            print(f"{i}. {category} ({weapon_count}种武器)")
        while True:
            try:
                choice = int(input(f"\n请选择类别 (1-{len(categories)}): "))
                if 1 <= choice <= len(categories):
                    selected_category = categories[choice-1]
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        print(f"\n{selected_category} 包含以下武器:")
        weapons = self.weapons_by_category[selected_category]
        for i, weapon in enumerate(weapons, 1):
            barrels_info = f" ({len(weapon['barrels'])}种枪管)" if len(weapon['barrels']) > 1 else ""
            display_caliber = weapon.get('original_caliber', weapon['caliber'])
            print(f"{i}. {weapon['name']} - {display_caliber}{barrels_info}")
        while True:
            try:
                weapon_choice = int(input(f"\n请选择武器 (1-{len(weapons)}): "))
                if 1 <= weapon_choice <= len(weapons):
                    selected_weapon = weapons[weapon_choice-1]
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        if self.debug:
            print(f"[DEBUG] 选中武器: {selected_weapon['name']} (ID={selected_weapon['id']})")
        return selected_weapon
    
    def select_barrel(self, weapon):
        print(f"\n{weapon['name']} 可用的枪管:")
        barrels = weapon['barrels']
        if len(barrels) == 1 and barrels[0]['name'] == "/":
            print("该武器只有默认枪管")
            return barrels[0]
        for i, barrel in enumerate(barrels, 1):
            power = barrel['firing_power']
            power_sign = "+" if power >= Decimal('0') else ""
            print(f"{i}. {barrel['name']} (击发威力: {power_sign}{power:.0f})")
        while True:
            try:
                choice = int(input(f"\n请选择枪管 (1-{len(barrels)}): "))
                if 1 <= choice <= len(barrels):
                    selected_barrel = barrels[choice-1]
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        if self.debug:
            print(f"[DEBUG] 选中枪管: {selected_barrel['name']} 威力={selected_barrel['firing_power']}")
        return selected_barrel
    
    def select_ammo(self, caliber):
        standardized_caliber = self.standardize_caliber(caliber)
        print(f"\n{caliber} 口径可用的弹药:")
        if self.debug:
            print(f"[DEBUG] 标准化口径: {standardized_caliber}")
        if standardized_caliber not in self.ammo_by_caliber:
            print(f"警告: 未找到标准化口径 '{standardized_caliber}' 的弹药数据！")
            print(f"可用口径: {', '.join(self.ammo_by_caliber.keys())}")
            matching_calibers = []
            for available_caliber in self.ammo_by_caliber.keys():
                if standardized_caliber in available_caliber or available_caliber in standardized_caliber:
                    matching_calibers.append(available_caliber)
            if matching_calibers:
                print(f"发现相似口径: {', '.join(matching_calibers)}")
                standardized_caliber = matching_calibers[0]
                print(f"自动选择: {standardized_caliber}")
            else:
                print("无法找到匹配的口径弹药！")
                return None
        ammo_list = self.ammo_by_caliber[standardized_caliber]
        for i, ammo in enumerate(ammo_list, 1):
            display_caliber = ammo.get('original_caliber', ammo['caliber'])
            print(f"{i}. {ammo['name']} - 伤害: {ammo['damage']:.0f}, 穿甲等级: {ammo['penetration_level']}, 护甲伤害: {ammo['armor_damage']:.1f}")
        while True:
            try:
                choice = int(input(f"\n请选择弹药 (1-{len(ammo_list)}): "))
                if 1 <= choice <= len(ammo_list):
                    selected_ammo = ammo_list[choice-1]
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        if self.debug:
            print(f"[DEBUG] 选中弹药: {selected_ammo['name']} (ID={selected_ammo['id']}, 子弹穿透={selected_ammo['bullet_penetration']})")
        return selected_ammo
    
    def select_armor(self):
        print("\n请选择护甲等级:")
        levels = sorted(self.armor_by_level.keys())
        for level in levels:
            armor_count = len(self.armor_by_level[level])
            print(f"{level}. 防护等级{level} ({armor_count}件)")
        while True:
            try:
                level = int(input(f"\n请选择防护等级 ({levels[0]}-{levels[-1]}): "))
                if level in levels:
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        print(f"\n防护等级 {level} 的护甲:")
        armors = self.armor_by_level[level]
        for i, armor in enumerate(armors, 1):
            material = armor.get('material', '未知材质')
            print(f"{i}. {armor['name']} - 材质: {material} - 耐久: {armor['durability']:.1f}")
        while True:
            try:
                choice = int(input(f"\n请选择护甲 (1-{len(armors)}): "))
                if 1 <= choice <= len(armors):
                    selected_armor = armors[choice-1]
                    break
                else:
                    print("输入无效，请重新选择。")
            except ValueError:
                print("请输入数字。")
        print("\n" + "=" * 50)
        print("护甲耐久设置")
        print("=" * 50)
        print(f"护甲: {selected_armor['name']}")
        print(f"材质: {selected_armor.get('material', '未知材质')}")
        print(f"初始耐久上限: {selected_armor['durability']:.1f}")
        while True:
            custom_input = input("\n是否设置自定义耐久？(Y/N, 输入N则使用初始耐久上限和耐久): ").strip().upper()
            if custom_input == 'Y':
                while True:
                    try:
                        max_input = input(f"\n请输入当前护甲上限 (最多一位小数, 0-{selected_armor['durability']:.1f}): ").strip()
                        if not max_input.replace('.', '').isdigit():
                            print("输入必须为数字，请重新输入。")
                            continue
                        if '.' in max_input:
                            decimal_parts = max_input.split('.')
                            if len(decimal_parts) > 2 or len(decimal_parts[1]) > 1:
                                print("格式错误，最多一位小数。")
                                continue
                        max_durability = Decimal(max_input)
                        if Decimal('0') <= max_durability <= selected_armor['durability']:
                            current_max_durability = max_durability
                            break
                        else:
                            print(f"当前上限必须在0-{selected_armor['durability']:.1f}之间。")
                    except Exception as e:
                        print(f"输入无效: {e}")
                while True:
                    try:
                        durability_input = input(f"\n请输入当前护甲耐久 (最多一位小数, 0-{current_max_durability:.1f}): ").strip()
                        if not durability_input.replace('.', '').isdigit():
                            print("输入必须为数字，请重新输入。")
                            continue
                        if '.' in durability_input:
                            decimal_parts = durability_input.split('.')
                            if len(decimal_parts) > 2 or len(decimal_parts[1]) > 1:
                                print("格式错误，最多一位小数。")
                                continue
                        current_durability = Decimal(durability_input)
                        if Decimal('0') <= current_durability <= current_max_durability:
                            break
                        else:
                            print(f"当前耐久必须在0-{current_max_durability:.1f}之间。")
                    except Exception as e:
                        print(f"输入无效: {e}")
                break
            elif custom_input == 'N':
                current_max_durability = selected_armor['durability']
                current_durability = selected_armor['durability']
                print(f"使用初始耐久上限: {current_max_durability:.1f}")
                print(f"使用初始耐久: {current_durability:.1f}")
                break
            else:
                print("输入无效，请输入 Y 或 N。")
        selected_armor['current_max_durability'] = current_max_durability
        selected_armor['current_durability'] = current_durability
        print(f"\n设置完成:")
        print(f"当前护甲上限: {current_max_durability:.1f}")
        print(f"当前护甲耐久: {current_durability:.1f}")
        if current_max_durability > Decimal('0'):
            durability_ratio = (current_durability / current_max_durability).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            print(f"耐久比例: {durability_ratio:.2%}")
        if self.debug:
            print(f"[DEBUG] 选中的护甲: {selected_armor['name']} (等级={selected_armor['level']}, 当前耐久={current_durability}, 上限={current_max_durability})")
        return selected_armor
    
    def select_armor_level(self):
        print("\n请选择护甲等级:")
        levels = sorted(self.armor_by_level.keys())
        for level in levels:
            print(f"{level}. 防护等级 {level}")
        while True:
            try:
                level = int(input(f"\n请输入护甲等级 ({levels[0]}-{levels[-1]}): "))
                if level in levels:
                    if self.debug:
                        print(f"[DEBUG] 选择护甲等级: {level}")
                    return level
                else:
                    print(f"输入无效，请选择 {levels[0]}-{levels[-1]} 之间的等级。")
            except ValueError:
                print("请输入数字。")
    
    def simulate_shooting(self):
        print("=" * 50)
        print(" ")
        print("=" * 50)
        weapon = self.select_weapon()
        barrel = self.select_barrel(weapon)
        display_caliber = weapon.get('original_caliber', weapon['caliber'])
        ammo = self.select_ammo(display_caliber)
        if not ammo:
            print("无法继续计算，请重新开始。")
            return True
        armor = self.select_armor()
        current_max_durability = armor.get('current_max_durability', armor['durability'])
        current_durability = armor.get('current_durability', armor['durability'])
        print("\n" + "=" * 50)
        print("模拟开始")
        print("=" * 50)
        weapon_display_caliber = weapon.get('original_caliber', weapon['caliber'])
        print(f"武器: {weapon['name']} ({weapon_display_caliber})")
        print(f"枪管: {barrel['name']}")
        ammo_display_caliber = ammo.get('original_caliber', ammo['caliber'])
        if ammo_display_caliber is None:
            ammo_display_caliber = ammo['caliber']
        print(f"弹药: {ammo['name']} ({ammo_display_caliber}, 穿透等级: {ammo['penetration_level']})")
        print(f"护甲: {armor['name']} - 材质: {armor.get('material', '未知材质')} (防护等级: {armor['level']})")
        print(f"初始耐久上限: {armor['durability']:.1f}")
        print(f"当前耐久上限: {current_max_durability:.1f}")
        print(f"当前耐久: {current_durability:.1f}")
        if current_max_durability > Decimal('0'):
            durability_ratio = (current_durability / current_max_durability).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            print(f"当前耐久比例: {durability_ratio:.2%}")
        print("=" * 50)
        round_num = 1
        first_probable_hit = None
        first_guaranteed_hit = None
        simulation_history = []
        while current_durability > Decimal('0'):
            result = self.calculate_damage(weapon, barrel, ammo, armor, current_durability, current_max_durability)
            simulation_history.append({
                'round': round_num,
                'penetration_type': result['penetration_type'],
                'damage': result['damage'],
                'remaining_durability': result['remaining_durability'],
                'armor_damage_dealt': result['armor_damage_dealt'],
                'can_penetrate': result['can_penetrate'],
                'current_durability_before': current_durability,
                'current_max_durability': current_max_durability
            })
            if first_probable_hit is None and result['penetration_type'] == "概率击穿护甲":
                first_probable_hit = round_num
            if first_guaranteed_hit is None and result['penetration_type'] == "必定击穿护甲":
                first_guaranteed_hit = round_num
            current_durability = result['remaining_durability']
            if result['is_destroyed']:
                break
            round_num += 1
        print("\n" + "=" * 50)
        print("击穿分析结果")
        print("=" * 50)
        if first_probable_hit is not None:
            print(f"概率击穿所需命中数：{first_probable_hit}")
        else:
            print("概率击穿所需命中数：未出现概率击穿")
        if first_guaranteed_hit is not None:
            print(f"稳定击穿所需命中数：{first_guaranteed_hit}")
        else:
            print("稳定击穿所需命中数：未出现稳定击穿")
        print(f"护甲在第{round_num}轮被彻底击毁！")
        print("\n" + "=" * 50)
        while True:
            show_details = input("是否显示详细计算过程？(Y/N): ").strip().upper()
            if show_details == 'Y':
                print("\n" + "=" * 50)
                print("详细计算过程")
                print("=" * 50)
                for i, result in enumerate(simulation_history):
                    if i < len(simulation_history) - 1:
                        print(f"\n第{result['round']}轮射击")
                        print(result['penetration_type'])
                        if result['can_penetrate'] and result['damage'] > Decimal('0'):
                            print(f"造成{result['damage']:.2f}伤害")
                            through_damage = result['damage'] * Decimal('0.88')
                            print(f"穿手伤害：{through_damage:.2f}")
                        print(f"剩余护甲耐久：{result['remaining_durability']:.1f}/{result['current_max_durability']:.1f}")
                        if result['current_max_durability'] > Decimal('0'):
                            durability_ratio = (result['remaining_durability'] / result['current_max_durability']).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
                            print(f"耐久比例：{durability_ratio:.2%}")
                    else:
                        print(f"\n第{result['round']}轮射击")
                        print(result['penetration_type'])
                        if result['can_penetrate'] and result['damage'] > Decimal('0'):
                            print(f"造成{result['damage']:.2f}伤害")
                            through_damage = result['damage'] * Decimal('0.88')
                            print(f"穿手伤害：{through_damage:.2f}")
                        print(f"剩余护甲耐久：{result['remaining_durability']:.1f}/{result['current_max_durability']:.1f}")
                        if result['current_max_durability'] > Decimal('0'):
                            durability_ratio = (result['remaining_durability'] / result['current_max_durability']).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
                            print(f"耐久比例：{durability_ratio:.2%}")
                        print("护甲被彻底击毁！")
                break
            elif show_details == 'N':
                break
            else:
                print("输入无效，请输入 Y 或 N。")
        armor['current_durability'] = armor['durability']
        print("\n" + "=" * 50)
        while True:
            answer = input("是否开始新一轮计算？(Y/N): ").strip().upper()
            if answer == 'Y':
                print("\n" + "=" * 50)
                print("开始新一轮计算")
                print("=" * 50)
                return True
            elif answer == 'N':
                return False
            else:
                print("输入无效，请输入 Y 或 N。")
    
    def damage_calculation_simulation(self):
        print("\n" + "=" * 50)
        print("伤害计算模拟")
        print("=" * 50)
        while True:
            weapon = self.select_weapon()
            barrel = self.select_barrel(weapon)
            display_caliber = weapon.get('original_caliber', weapon['caliber'])
            ammo = self.select_ammo(display_caliber)
            if not ammo:
                print("弹药选择失败，返回主菜单。")
                return
            armor_level = self.select_armor_level()
            total_damage = ammo['damage'] + barrel['firing_power']
            over_range_total = total_damage * Decimal('0.7')
            k_value = ammo['bullet_penetration'] - Decimal(str(armor_level)) * Decimal('10')
            penetration_percentage = self.calculate_penetration_percentage(k_value)
            armor_penetration_damage = (total_damage * penetration_percentage).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            through_hand_penetration = (armor_penetration_damage * Decimal('0.88')).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            over_range_penetration = (armor_penetration_damage * Decimal('0.7')).quantize(Decimal('0.0001'), rounding=ROUND_HALF_UP)
            if self.debug:
                print("\n[DEBUG] 伤害计算详细信息:")
                print(f"  弹药伤害: {ammo['damage']}")
                print(f"  枪管威力: {barrel['firing_power']}")
                print(f"  总伤害: {total_damage}")
                print(f"  护甲等级: {armor_level}")
                print(f"  子弹穿透值: {ammo['bullet_penetration']}")
                print(f"  K值: {k_value}")
                print(f"  穿透百分比: {penetration_percentage:.4f}")
                print(f"  穿甲伤害(未取整): {total_damage * penetration_percentage}")
                print(f"  最终穿甲伤害: {armor_penetration_damage}")
            print("\n" + "=" * 50)
            print("计算结果")
            print("=" * 50)
            print(f"总伤害：{total_damage:.2f}")
            print(f"超射程伤害：{over_range_total:.2f}")
            print(f"穿甲伤害：{armor_penetration_damage:.2f}")
            print(f"穿手穿甲伤害：{through_hand_penetration:.2f}")
            print(f"超射程穿甲伤害：{over_range_penetration:.2f}")
            print("=" * 50)
            while True:
                answer = input("\n是否进行新一轮伤害计算？(Y/N): ").strip().upper()
                if answer == 'Y':
                    print("\n" + "=" * 50)
                    print("开始新一轮伤害计算")
                    print("=" * 50)
                    break
                elif answer == 'N':
                    print("返回主菜单。")
                    return
                else:
                    print("输入无效，请输入 Y 或 N。")
    
    def run(self):
        print("=" * 50)
        print("护甲系统模拟器")
        print("=" * 50)
        print("击穿规则:")
        print("1. 子弹穿透等级≥护甲等级：必定击穿护甲")
        print("2. 子弹穿透等级＜护甲等级：根据护甲耐久计算")
        print("3. 护甲耐久减半时，必定能被更低一级的子弹击穿")
        print("=" * 50)
        print("耐久比例计算：当前耐久/当前上限")
        print("如未设置自定义耐久，则使用初始耐久上限")
        print("=" * 50)
        while True:
            print("\n请选择功能：")
            print("1. 护甲损毁计算模拟")
            print("2. 伤害计算模拟")
            print("0. 退出")
            choice = input("请输入选项：").strip()
            if choice == '1':
                while True:
                    cont = self.simulate_shooting()
                    if not cont:
                        break
            elif choice == '2':
                self.damage_calculation_simulation()
            elif choice == '0':
                print("退出程序。")
                break
            else:
                print("输入无效，请重新选择。")

def main():
    try:
        required_files = ['穿透数据整理结果.xlsx', '武器数据.xlsx', '子弹数据.xlsx', '护甲信息.xlsx']
        missing_files = [f for f in required_files if not os.path.exists(f)]
        if missing_files:
            print("警告: 以下文件不存在:")
            for f in missing_files:
                print(f"  - {f}")
            print("请确保所有数据文件在当前目录下。")
            input("按 Enter 键退出...")
            return
        calculator = ArmorPenetrationCalculator()
        calculator.run()
    except Exception as e:
        print("\n" + "=" * 50)
        print("程序运行出错！错误信息如下：")
        print("=" * 50)
        import traceback
        traceback.print_exc()
        print("=" * 50)
        input("按 Enter 键退出...")

if __name__ == "__main__":
    main()
